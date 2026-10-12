import os
import re
import sys
import json
import math
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv()

MODEL = "nvidia/nemotron-3-ultra-550b-a55b:free"   
MAX_STEPS = 12

# =====================================================================
# SYSTEM PROMPT (Step 3) + WORKFLOW (Step 5)
# =====================================================================
SYSTEM_PROMPT = """
You are a senior web application security researcher who writes
professional vulnerability reports for bug bounty programs and
penetration tests. All testing is authorized and in scope.

REPORT STRUCTURE (always use these sections, in this order)
1. Title (vulnerability type + affected component)
2. Summary (2-3 sentences, plain language)
3. OWASP Category / CWE
4. Severity (CVSS vector, score, reasoning)
5. Affected Asset
6. Steps to Reproduce (numbered, exact, repeatable)
7. Proof of Concept (requests/responses supplied by the user)
8. Impact (realistic, specific, no exaggeration)
9. Remediation (actionable fixes)
10. References

RULES
- Use ONLY facts the user provided. Never invent URLs, parameters,
  payloads, responses or impact.
- If information is missing and not essential, write
  [NEEDS INFO: what is missing] instead of guessing.
- Never calculate CVSS yourself. Use the calculate_cvss tool.
- Never rely on memory for OWASP categories or CWE IDs. Use lookup_owasp
  and cite the edition it returns.
- Do not exaggerate severity. Keep the tone professional and concise.

WORKFLOW
1. Call lookup_owasp to get the category and CWE.
2. If essential facts are missing (for example the exact request, the
   affected endpoint, or what data was exposed), call ask_user. Ask at
   most 3 questions in total.
3. Choose CVSS metrics from the facts, then call calculate_cvss.
4. Write the full draft report.
5. Call quality_check with the draft. If it fails, fix and check again.
6. When it passes, call save_report once.
7. Finish by telling the user the file names and any [NEEDS INFO] items.
"""

# =====================================================================
# STEP 6: OWASP knowledge (reads owasp_data.json)
# =====================================================================
def lookup_owasp(query):
    with open("owasp_data.json", encoding="utf-8") as f:
        data = json.load(f)
    q = query.lower()
    cats = [c for c in data["categories"]
            if q in c["name"].lower() or any(k in q or q in k for k in c["keywords"])]
    cwes = {k: v for k, v in data["cwes"].items() if k.lower() in q or q in k.lower()}
    if not cats and not cwes:
        return f"No match for '{query}'. Edition on file: {data['edition']}. Try another term."
    lines = [f"Edition on file: {data['edition']}"]
    for c in cats:
        lines.append(f"Category: {c['id']} {c['name']}")
    for k, v in cwes.items():
        lines.append(f"CWE for {k}: {v}")
    return "\n".join(lines)

# =====================================================================
# STEP 7: CVSS v3.1 base score (done by code, not by the AI)
# =====================================================================
def _roundup(x):
    i = round(x * 100000)
    if i % 10000 == 0:
        return i / 100000.0
    return (math.floor(i / 10000) + 1) / 10.0

def calculate_cvss(AV, AC, PR, UI, S, C, I, A):
    changed = S == "C"
    av = {"N": 0.85, "A": 0.62, "L": 0.55, "P": 0.2}[AV]
    ac = {"L": 0.77, "H": 0.44}[AC]
    pr = {"N": 0.85, "L": 0.68 if changed else 0.62, "H": 0.5 if changed else 0.27}[PR]
    ui = {"N": 0.85, "R": 0.62}[UI]
    cia = {"H": 0.56, "L": 0.22, "N": 0.0}
    iss = 1 - (1 - cia[C]) * (1 - cia[I]) * (1 - cia[A])
    if changed:
        impact = 7.52 * (iss - 0.029) - 3.25 * (iss - 0.02) ** 15
    else:
        impact = 6.42 * iss
    expl = 8.22 * av * ac * pr * ui
    if impact <= 0:
        score = 0.0
    elif changed:
        score = _roundup(min(1.08 * (impact + expl), 10))
    else:
        score = _roundup(min(impact + expl, 10))
    if score == 0:
        sev = "None"
    elif score < 4:
        sev = "Low"
    elif score < 7:
        sev = "Medium"
    elif score < 9:
        sev = "High"
    else:
        sev = "Critical"
    vector = f"CVSS:3.1/AV:{AV}/AC:{AC}/PR:{PR}/UI:{UI}/S:{S}/C:{C}/I:{I}/A:{A}"
    return f"Vector: {vector} | Score: {score} | Severity: {sev}"

# =====================================================================
# STEP 8: ask the human a question
# =====================================================================
def ask_user(question):
    print(f"\nAGENT ASKS: {question}")
    return input("YOUR ANSWER: ").strip() or "(no answer given)"

# =====================================================================
# STEP 9: quality checker (plain code, no AI)
# =====================================================================
REQUIRED_SECTIONS = [
    "Summary", "OWASP", "Severity", "Affected Asset",
    "Steps to Reproduce", "Proof of Concept",
    "Impact", "Remediation", "References",
]
VAGUE_PHRASES = [
    "could potentially", "might be vulnerable", "may be vulnerable",
    "very dangerous", "extremely critical", "catastrophic",
    "hackers can do anything",
]

def quality_check(content):
    problems = []
    low = content.lower()
    missing = [s for s in REQUIRED_SECTIONS if s.lower() not in low]
    if missing:
        problems.append("Missing sections: " + ", ".join(missing))
    if not re.search(r"CVSS:3\.1/AV:", content):
        problems.append("No CVSS vector. Use the calculate_cvss output.")
    if not re.search(r"CWE-\d+", content):
        problems.append("No CWE ID.")
    if not re.search(r"A\d{2}:\d{4}", content):
        problems.append("No OWASP category ID (like A01:2021).")
    if len(re.findall(r"(?m)^\s*\d+\.\s", content)) < 3:
        problems.append("Steps to Reproduce need at least 3 numbered steps.")
    found = [p for p in VAGUE_PHRASES if p in low]
    if found:
        problems.append("Vague or exaggerated wording: " + ", ".join(found))
    if problems:
        return "FAIL. " + " | ".join(problems)
    return f"PASS. Open [NEEDS INFO] items: {content.count('[NEEDS INFO')}"

# =====================================================================
# STEP 10: save as Markdown + Word
# =====================================================================
def _md_to_docx(md_text, path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    in_code = False
    for line in md_text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], 1)
        elif re.match(r"^\s*[-*]\s", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s", "", line).replace("**", ""), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.replace("**", ""))
    doc.save(path)

def save_report(filename, content):
    os.makedirs("reports", exist_ok=True)
    base = os.path.splitext(os.path.basename(filename))[0]   # blocks ../ tricks
    md_path = os.path.join("reports", base + ".md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    msg = f"Saved {md_path}"
    try:
        docx_path = os.path.join("reports", base + ".docx")
        _md_to_docx(content, docx_path)
        msg += f" and {docx_path}"
    except ImportError:
        msg += " (Word export skipped: run pip install python-docx)"
    return msg

# =====================================================================
# STEP 11: safety helpers
# =====================================================================
def redact(text):
    """Hide emails, phone-like numbers and tokens BEFORE sending text to the AI."""
    text, n1 = re.subn(r"[\w.+-]+@[\w-]+\.[\w.-]+", "[REDACTED-EMAIL]", text)
    text, n2 = re.subn(r"\+?\d[\d\s().-]{8,}\d", "[REDACTED-NUMBER]", text)
    text, n3 = re.subn(r"(?i)(bearer|token|api[_-]?key|authorization)([:=\s]+)[A-Za-z0-9._\-]{16,}",
                       r"\1\2[REDACTED-TOKEN]", text)
    print(f"Redacted: {n1} emails, {n2} number-like values, {n3} tokens")
    return text

# =====================================================================
# Tool registry (what the model sees + what your code runs)
# =====================================================================
def _tool(name, desc, props, required):
    return {"type": "function", "function": {
        "name": name, "description": desc,
        "parameters": {"type": "object", "properties": props, "required": required}}}

def _enum(desc, values):
    return {"type": "string", "enum": values, "description": desc}

TOOL_FUNCTIONS = {
    "lookup_owasp": lookup_owasp,
    "calculate_cvss": calculate_cvss,
    "ask_user": ask_user,
    "quality_check": quality_check,
    "save_report": save_report,
}

TOOLS = [
    _tool("lookup_owasp",
          "Find the OWASP category and CWE for a vulnerability type (e.g. 'IDOR', 'XSS').",
          {"query": {"type": "string"}}, ["query"]),
    _tool("calculate_cvss",
          "Calculate the CVSS v3.1 base score from metrics. Always use this for severity.",
          {"AV": _enum("Attack Vector", ["N", "A", "L", "P"]),
           "AC": _enum("Attack Complexity", ["L", "H"]),
           "PR": _enum("Privileges Required", ["N", "L", "H"]),
           "UI": _enum("User Interaction", ["N", "R"]),
           "S": _enum("Scope", ["U", "C"]),
           "C": _enum("Confidentiality impact", ["H", "L", "N"]),
           "I": _enum("Integrity impact", ["H", "L", "N"]),
           "A": _enum("Availability impact", ["H", "L", "N"])},
          ["AV", "AC", "PR", "UI", "S", "C", "I", "A"]),
    _tool("ask_user",
          "Ask the human one short question when essential information is missing.",
          {"question": {"type": "string"}}, ["question"]),
    _tool("quality_check",
          "Check a draft report for structure and quality problems. Call before saving.",
          {"content": {"type": "string", "description": "Full draft in Markdown"}}, ["content"]),
    _tool("save_report",
          "Save the final report (Markdown + Word). Only call after quality_check passes.",
          {"filename": {"type": "string", "description": "Short name, e.g. idor-orders-api"},
           "content": {"type": "string", "description": "Full report in Markdown"}},
          ["filename", "content"]),
]

# =====================================================================
# THE LOOP (Step 5)
# =====================================================================
def run_agent(client, task):
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": task},
    ]
    for step in range(1, MAX_STEPS + 1):
        response = client.chat.completions.create(model=MODEL, messages=messages, tools=TOOLS)
        if response.choices is None:
            print("API error:", response)
            return
        msg = response.choices[0].message
        messages.append(msg)
        if not msg.tool_calls:
            print("\nFINAL ANSWER:\n", msg.content)
            return
        for call in msg.tool_calls:
            name = call.function.name
            print(f"[step {step}] tool requested: {name}")
            try:
                args = json.loads(call.function.arguments)
                result = TOOL_FUNCTIONS[name](**args)
            except Exception as e:
                result = f"Tool error: {e}"
            print(f"           result: {result[:200]}")
            messages.append({"role": "tool", "tool_call_id": call.id, "content": result})
    print("Stopped: reached MAX_STEPS without finishing.")

SAMPLE_FINDING = """
Target: shop.example.com (my own test lab)
I logged in as user A. Requested GET /api/orders/123 (my order).
Changed 123 to 124 and got another user's order with name,
address and phone number. No authorization error.
"""

def main():
    answer = input("Do you have authorization to test this target "
                   "(your own lab or an in-scope program)? y/n: ")
    if answer.strip().lower() != "y":
        print("Stopping. Only report on authorized, in-scope testing.")
        return
    if len(sys.argv) > 1:
        with open(sys.argv[1], encoding="utf-8") as f:
            finding = f.read()
    else:
        finding = SAMPLE_FINDING
    finding = redact(finding)
    client = OpenAI(
        base_url="https://openrouter.ai/api/v1",
        api_key=os.environ["OPEN_ROUTER_API_KEY"],
    )
    run_agent(client, f"Write a report for this finding:\n{finding}")

if __name__ == "__main__":
    main()
